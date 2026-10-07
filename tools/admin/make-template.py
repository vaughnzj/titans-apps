#!/usr/bin/env python3
"""Build the Legend Titans roster entry template.

The column headers are the contract: the Roster Manager maps them by name, in any
order, so these must match what its importer recognises — Last Name / First Name /
Class / Throws /
Bats / Jersey. Everything else on the sheet is there to keep the data clean before
it ever reaches the apps.
"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule

# Run with no arguments to build the blank template. Pass --names FILE.tsv (last,
# first, note per line) and --out PATH to build a filled-in roster from it: the
# example rows are left out, since real data is already there to look at.
import argparse
ap = argparse.ArgumentParser()
ap.add_argument("--out", default="/home/claude/Titans-Roster-Template.xlsx")
ap.add_argument("--names", default=None, help="TSV: last, first, note")
ap.add_argument("--title", default="LEGEND TITANS \u2014 PROGRAM ROSTER")
A = ap.parse_args()

PLAYERS = []
if A.names:
    for line in io.open(A.names, encoding="utf-8").read().splitlines():
        if not line.strip():
            continue
        p = (line.split("\t") + ["", "", ""])[:3]
        PLAYERS.append(p)

OUT = A.out
# Room for the whole program with space to spare. With names already in, keep at
# least 25 blank rows under them so adding a late arrival needs no resizing.
ROWS = max(80, len(PLAYERS) + 25)

NAVY   = "16305C"
NAVY_2 = "1D3F76"
RUST   = "C85C1C"
LINE   = "D2D9E6"
HEAD_T = "FFFFFF"
INPUT  = "FFFFCC"              # the fill that means "you type here"
GREY   = "F4F6FA"
BAD    = "FAE5E1"

thin = Side(style="thin", color=LINE)
box  = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()

# ───────────────────────────── Roster ─────────────────────────────
ws = wb.active
ws.title = "Roster"

ws["A1"] = A.title
ws["A1"].font = Font(name="Arial", size=15, bold=True, color=NAVY)
ws["A2"] = ("Type into the shaded cells. Save, then upload this file in the Roster Manager "
            "to release a new roster.js.")
ws["A2"].font = Font(name="Arial", size=10, italic=True, color="5A6781")
ws.merge_cells("A1:G1")
ws.merge_cells("A2:G2")

# Last and First are SEPARATE columns because that is how a program roster is kept,
# and because one combined field is what let a senior Dieker and a freshman Dieker
# import as one player. The pair is the identity; the apps derive what to print.
HEADERS = ["Last Name", "First Name", "Class", "Throws", "Bats", "Jersey",
           "Notes (not imported)"]
HROW = 4
for i, h in enumerate(HEADERS, start=1):
    c = ws.cell(row=HROW, column=i, value=h)
    c.font = Font(name="Arial", size=11, bold=True, color=HEAD_T)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = box
ws.row_dimensions[HROW].height = 22

NOTES = {
    "A": "Last name only. Put the first name in the next column.\n\n"
         "Last + First together are the identity. Two players may share a last name — "
         "that is handled, they get separate IDs, and the apps show \"Dieker, T\" and "
         "\"Dieker, J\" automatically.",
    "B": "First name. Worth filling in even when nobody shares the last name:\n\n"
         "it is the only thing available to tell two players apart if a brother or "
         "another same-surname kid joins later. The apps show just the last name until "
         "that happens.",
    "C": "Fr, So, Jr or Sr. Pick from the list.\n\n"
         "This is what makes a peer comparison possible: a sophomore measured against "
         "sophomores says far more than one measured against the whole program.\n\n"
         "After the season ends, the Roster Manager moves everyone up a class and takes "
         "the graduating seniors off the roster.",
    "D": "Throwing hand: R or L.",
    "E": "Batting side: R, L or S (switch).\n\n"
         "Separate from Throws on purpose — plenty of kids throw right and hit left.",
    "F": "Optional. The Offense Chart uses it; the Bullpen Chart ignores it.",
    "G": "Yours to use. This column is NOT imported — nothing here reaches the apps.",
}
for col, text in NOTES.items():
    ws[f"{col}{HROW}"].comment = Comment(text, "Roster Manager")

# Two example rows, and they share a last name on purpose: that is the case worth
# showing, because it is the one people assume will break.
ex = HROW + 1
EXAMPLES = [
    ("Hasche", "Brock", "Jr", "R", "L", "17",
     "EXAMPLE ROW — delete all three before you upload"),
    ("Dieker", "Tom",   "Sr", "R", "R", "4",
     "EXAMPLE — two Diekers is fine. They get separate IDs, and the apps "
     "show \"Dieker, T\" and \"Dieker, J\""),
    ("Dieker", "Jack",  "Fr", "L", "L", "",
     "EXAMPLE — delete these three rows"),
]
if PLAYERS:
    EXAMPLES = []              # real names are a better example than invented ones
for n, vals in enumerate(EXAMPLES):
    for i, v in enumerate(vals, start=1):
        ws.cell(row=ex + n, column=i, value=v)
    for i in range(1, 8):
        c = ws.cell(row=ex + n, column=i)
        c.font = Font(name="Arial", size=11, italic=True, color="8793A9")
        c.border = box
        c.alignment = Alignment(horizontal="left" if i in (1, 2, 7) else "center")
EXLAST = ex + len(EXAMPLES) - 1 if EXAMPLES else HROW

FIRST = EXLAST + 1
LAST = HROW + ROWS
for r in range(FIRST, LAST + 1):
    for i in range(1, 8):
        c = ws.cell(row=r, column=i)
        c.font = Font(name="Arial", size=11)
        c.border = box
        c.alignment = Alignment(horizontal="left" if i in (1, 2, 7) else "center")
        if i <= 6:
            c.fill = PatternFill("solid", fgColor=INPUT)

# The names, if any. Class / Throws / Bats / Jersey are left blank on purpose -
# Jim fills those; nothing here guesses at a class year from a team sheet, because
# Varsity and JV are mixed classes and a wrong class quietly poisons every peer
# comparison it touches.
for n, (plast, pfirst, pnote) in enumerate(PLAYERS):
    r = FIRST + n
    ws.cell(row=r, column=1, value=plast)
    ws.cell(row=r, column=2, value=pfirst)
    if pnote:
        ws.cell(row=r, column=7, value=pnote)

# Jersey stays text — a leading zero on "07" must survive.
for r in range(ex, LAST + 1):
    ws.cell(row=r, column=6).number_format = "@"

for col, w in zip("ABCDEFG", [20, 16, 9, 10, 9, 10, 40]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = f"A{FIRST}"

# ───────────────────────── dropdowns ─────────────────────────
# Pointed at the Lists sheet so the allowed values are visible and editable rather
# than buried in the validation dialog.
dv_cls = DataValidation(type="list", formula1="=Lists!$A$2:$A$5", allow_blank=True,
                        showDropDown=False, errorStyle="stop",
                        errorTitle="Not a class",
                        error="Pick Fr, So, Jr or Sr.",
                        promptTitle="Class", prompt="Fr, So, Jr or Sr")
dv_thr = DataValidation(type="list", formula1="=Lists!$B$2:$B$3", allow_blank=True,
                        showDropDown=False, errorStyle="stop",
                        errorTitle="Not a hand", error="Throwing hand is R or L.")
dv_bat = DataValidation(type="list", formula1="=Lists!$C$2:$C$4", allow_blank=True,
                        showDropDown=False, errorStyle="stop",
                        errorTitle="Not a side", error="Batting side is R, L or S.")
for dv, col in ((dv_cls, "C"), (dv_thr, "D"), (dv_bat, "E")):
    ws.add_data_validation(dv)
    dv.add(f"{col}{FIRST}:{col}{LAST}")

# ──────────────────── the duplicate guard ────────────────────
# The blocking problem is the same FIRST and LAST name — those are indistinguishable
# to a coach and to every spreadsheet downstream, and the Roster Manager refuses to
# release them. Two players sharing only a last name are FINE and must not be
# flagged; SUMPRODUCT counts the pair so they aren't.
dupe = FormulaRule(
    formula=[f'AND($A{FIRST}<>"",'
             f'SUMPRODUCT(($A${FIRST}:$A${LAST}=$A{FIRST})*($B${FIRST}:$B${LAST}=$B{FIRST}))>1)'],
    fill=PatternFill("solid", fgColor=BAD), stopIfTrue=False)
ws.conditional_formatting.add(f"A{FIRST}:B{LAST}", dupe)

# a last name with no first name: not an error, but worth seeing
nofirst = FormulaRule(
    formula=[f'AND($A{FIRST}<>"",$B{FIRST}="")'],
    fill=PatternFill("solid", fgColor="FAF0DA"), stopIfTrue=False)
ws.conditional_formatting.add(f"B{FIRST}:B{LAST}", nofirst)

# a name typed with no class is the other thing worth seeing early
nocls = FormulaRule(
    formula=[f'AND($A{FIRST}<>"",$C{FIRST}="")'],
    fill=PatternFill("solid", fgColor="FAF0DA"), stopIfTrue=False)
ws.conditional_formatting.add(f"C{FIRST}:C{LAST}", nocls)

# ───────────────────────── Check sheet ─────────────────────────
chk = wb.create_sheet("Check")
chk["A1"] = "BEFORE YOU UPLOAD"
chk["A1"].font = Font(name="Arial", size=14, bold=True, color=NAVY)
chk.merge_cells("A1:C1")

rows = [
    ("Players entered",
     f'=COUNTA(Roster!$A${FIRST}:$A${LAST})',
     "Every row with a name."),
    ("Same first AND last name",
     f'=SUMPRODUCT((Roster!$A${FIRST}:$A${LAST}<>"")*'
     f'(COUNTIFS(Roster!$A${FIRST}:$A${LAST},Roster!$A${FIRST}:$A${LAST},'
     f'Roster!$B${FIRST}:$B${LAST},Roster!$B${FIRST}:$B${LAST})>1))',
     "Must be 0. Two players nobody can tell apart — not a coach, not a spreadsheet. "
     "The Roster Manager refuses to release them."),
    ("Shared last names",
     f'=SUMPRODUCT((Roster!$A${FIRST}:$A${LAST}<>"")*'
     f'(COUNTIF(Roster!$A${FIRST}:$A${LAST},Roster!$A${FIRST}:$A${LAST})>1))',
     "Not a problem — this is fine. Each gets his own ID, and the apps show "
     "\"Dieker, T\" and \"Dieker, J\" instead of just \"Dieker\". Shown so the number "
     "isn't a surprise."),
    ("Missing a first name",
     f'=SUMPRODUCT((Roster!$A${FIRST}:$A${LAST}<>"")*(Roster!$B${FIRST}:$B${LAST}=""))',
     "Not blocking. But if a brother joins next year there is nothing to tell them "
     "apart with."),
    ("Missing a class",
     f'=SUMPRODUCT((Roster!$A${FIRST}:$A${LAST}<>"")*(Roster!$C${FIRST}:$C${LAST}=""))',
     "They still chart fine — they just drop out of any comparison against their own class."),
    ("Missing a throwing hand",
     f'=SUMPRODUCT((Roster!$A${FIRST}:$A${LAST}<>"")*(Roster!$D${FIRST}:$D${LAST}=""))',
     "Defaults to R on import if left blank."),
    ("Example rows still there",
     f'=COUNTIF(Roster!$G${ex}:$G${LAST},"EXAMPLE*")',
     "Delete all three, or Hasche and two Diekers join your roster."),
]
r = 3
for label, formula, why in rows:
    chk.cell(row=r, column=1, value=label).font = Font(name="Arial", size=11, bold=True)
    v = chk.cell(row=r, column=2, value=formula)
    v.font = Font(name="Arial", size=11, bold=True)
    v.alignment = Alignment(horizontal="center")
    v.number_format = "0"
    v.border = box
    chk.cell(row=r, column=3, value=why).font = Font(name="Arial", size=10, color="5A6781")
    chk.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    chk.row_dimensions[r].height = 30
    r += 1

# anything that must be zero goes red when it isn't
# Only two rows MUST be zero: the same-name duplicate (row 4) and the leftover
# example rows (row 9). Shared last names and missing first names are not errors
# and deliberately get no red.
must_be_zero = PatternFill("solid", fgColor=BAD)
ok_fill = PatternFill("solid", fgColor="E2F3EA")
for rr in (4, 9):
    chk.conditional_formatting.add(
        f"B{rr}",
        FormulaRule(formula=[f"$B${rr}>0"], fill=must_be_zero, stopIfTrue=False))
    chk.conditional_formatting.add(
        f"B{rr}",
        FormulaRule(formula=[f"$B${rr}=0"], fill=ok_fill, stopIfTrue=False))

chk["A11"] = "Counts by class"
chk["A11"].font = Font(name="Arial", size=12, bold=True, color=NAVY)
r = 12
for code, name in (("Fr", "Freshmen"), ("So", "Sophomores"), ("Jr", "Juniors"), ("Sr", "Seniors")):
    chk.cell(row=r, column=1, value=name).font = Font(name="Arial", size=11)
    c = chk.cell(row=r, column=2,
                 value=f'=COUNTIF(Roster!$C${FIRST}:$C${LAST},"{code}")')
    c.font = Font(name="Arial", size=11)
    c.alignment = Alignment(horizontal="center")
    c.number_format = "0"
    c.border = box
    r += 1

chk["A17"] = "What to do with this file"
chk["A17"].font = Font(name="Arial", size=12, bold=True, color=NAVY)
steps = [
    "1.  Fill in the Roster tab. Delete the three example rows.",
    "2.  Check that the two rows above marked \"Must be 0\" both read 0.",
    "3.  Save this file.",
    "4.  Open the Roster Manager. Load the released roster.js first, if there is one.",
    "5.  Upload this file. Names already known keep their IDs; new names get new ones.",
    "6.  Download roster.js, drop it in the repo, bump the cache name in each app's sw.js.",
]
r = 18
for s in steps:
    chk.cell(row=r, column=1, value=s).font = Font(name="Arial", size=10.5)
    chk.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    r += 1

chk["A26"] = ("Column headers are the contract. The Roster Manager matches them by name in any "
              "order, so don't rename Last Name, First Name, Class, Throws, Bats or Jersey. "
              "Extra columns are ignored, and title rows above the headers are skipped.")
chk["A26"].font = Font(name="Arial", size=10, italic=True, color="5A6781")
chk["A26"].alignment = Alignment(wrap_text=True, vertical="top")
chk.merge_cells("A26:C28")

for col, w in zip("ABC", [26, 12, 62]):
    chk.column_dimensions[col].width = w

# ───────────────────────── Lists sheet ─────────────────────────
lst = wb.create_sheet("Lists")
lst["A1"] = "Class"; lst["B1"] = "Throws"; lst["C1"] = "Bats"
for cell in ("A1", "B1", "C1"):
    lst[cell].font = Font(name="Arial", size=11, bold=True, color=HEAD_T)
    lst[cell].fill = PatternFill("solid", fgColor=NAVY_2)
for i, v in enumerate(["Fr", "So", "Jr", "Sr"], start=2):
    lst.cell(row=i, column=1, value=v).font = Font(name="Arial", size=11)
for i, v in enumerate(["R", "L"], start=2):
    lst.cell(row=i, column=2, value=v).font = Font(name="Arial", size=11)
for i, v in enumerate(["R", "L", "S"], start=2):
    lst.cell(row=i, column=3, value=v).font = Font(name="Arial", size=11)
lst["E1"] = "These feed the dropdowns on the Roster tab. The apps only understand these values."
lst["E1"].font = Font(name="Arial", size=10, italic=True, color="5A6781")
for col, w in zip("ABCDE", [8, 9, 8, 3, 70]):
    lst.column_dimensions[col].width = w

wb.save(OUT)
print("wrote " + OUT)
